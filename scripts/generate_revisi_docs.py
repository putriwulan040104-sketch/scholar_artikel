# -*- coding: utf-8 -*-
"""Generator 2 dokumen revisi laporan TA berdasarkan 5 poin revisi dosen.
  1) Revisi_Narasi_5Poin.docx   -> narasi copy-paste ([SISIPKAN]/[GANTI])
  2) BabIV_Gabungan_Revisi.docx -> Bab IV utuh hasil + pembahasan revisi
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor

OUT_DIR = r"D:\Tugas Akhir\paperci_artikel"

# ===========================================================================
# DATA 33 GAMBAR KODE (Gambar 4.1 s.d. 4.33)
# setiap dict: t(judul), i(input), p(proses), o(output), j(kalimat penutup)
# ===========================================================================
F = [
 dict(t="run_web_ir_pipeline() pada ir_pipeline_service.py",
      i="Keyword pencarian yang diketik pengguna pada halaman web, diteruskan ke backend Flask melalui endpoint /search-progress sebagai parameter query.",
      p="Fungsi run_web_ir_pipeline() dijalankan dalam thread terpisah menggunakan Server-Sent Events (SSE) agar proses yang panjang tidak memblokir server. Kode ini mengoordinasikan seluruh alur: scraping, validasi metadata, penghapusan duplikat, hingga penyusunan indeks untuk keperluan Information Retrieval.",
      o="Kemajuan proses per tahap yang dilaporkan secara real-time ke frontend, serta dataset artikel yang siap diproses pada tahap berikutnya.",
      j="Tujuan dari blok kode ini adalah mengorkestrasi seluruh pipeline web Information Retrieval agar proses berjalan asinkron sekaligus memantau progres setiap tahap."),

 dict(t="pengambilan data berdasarkan keyword (_snapshot_scholar_results)",
      i="Hasil pencarian Google Scholar berupa kumpulan elemen artikel yang diperoleh berdasarkan keyword.",
      p="Fungsi _snapshot_scholar_results() mengambil metadata penting dari setiap artikel: judul, URL artikel, nama penulis, tahun publikasi, sumber (publisher/jurnal), abstrak (jika tersedia), dan DOI yang terdeteksi dari Google Scholar.",
      o="Snapshot metadata lengkap per artikel yang menjadi bahan temporary dataset sebelum divalidasi lanjut.",
      j="Tujuan dari blok kode ini adalah memotret seluruh metadata hasil scraping agar setiap artikel memiliki informasi yang cukup sebelum diproses lebih lanjut."),

 dict(t="pengambilan data berdasarkan keyword (pre-filtering metadata)",
      i="result_data hasil snapshot metadata dari Google Scholar.",
      p="Sistem melakukan validasi awal (pre-filtering): judul harus tersedia dan memiliki panjang memadai, tahun berada dalam rentang yang ditentukan, serta sumber tersedia. Artikel yang tidak memenuhi kriteria dilewati.",
      o="Temporary dataset yang telah lolos filter awal dan bebas dari data yang tidak layak.",
      j="Tujuan dari blok kode ini adalah menyaring metadata yang tidak valid di tahap awal guna menghindari pemrosesan artikel cacat lebih lanjut."),

 dict(t="validasi similarity checking artikel (1)",
      i="Field title dan abstract dari setiap artikel pada temporary dataset.",
      p="Judul dan abstrak digabung menjadi dokumen pembanding, diubah menjadi bobot TF-IDF menggunakan TfidfVectorizer, lalu dihitung cosine similarity antar artikel. Artikel dengan nilai kemiripan di atas ambang dikelompokkan dalam satu grup.",
      o="Grup artikel yang memiliki kemiripan tinggi (calon duplikat).",
      j="Tujuan dari blok kode ini adalah mendeteksi artikel yang isinya sangat serupa agar dapat direduksi menjadi satu kandidat pada tahap pemilihan sumber."),

 dict(t="validasi similarity checking artikel (2)",
      i="Grup artikel yang mirip beserta nilai prioritas sumber masing-masing artikel.",
      p="Fungsi get_official_source_priority() memberikan prioritas sumber (official publisher, official journal, institutional repository, mirror platform), lalu sistem memilih satu artikel dengan prioritas tertinggi untuk dipertahankan.",
      o="Satu artikel terbaik dari setiap grup artikel yang mirip.",
      j="Tujuan dari blok kode ini adalah menyeleksi artikel dari sumber resmi agar dataset lebih bersih dari artikel berulang tanpa kehilangan kualitas sumbernya."),

 dict(t="validasi DOI (1)",
      i="Metadata artikel berupa judul dan nama penulis hasil scraping.",
      p="Sistem mencoba mengambil DOI langsung dari halaman artikel melalui fetch_doi_from_doi_org(); bila belum ditemukan, dilakukan pencarian cadangan melalui API CrossRef dengan judul dan penulis (fetch_doi_from_crossref()).",
      o="DOI yang valid untuk artikel (atau status bahwa DOI tidak tersedia).",
      j="Tujuan dari blok kode ini adalah memastikan setiap artikel memiliki DOI yang valid sebelum dimasukkan ke final dataset sebagai identitas unik."),

 dict(t="validasi DOI (2) melalui CrossRef",
      i="Judul artikel sebagai parameter utama dan nama belakang penulis pertama sebagai parameter tambahan.",
      p="Sistem mengirim permintaan ke API CrossRef dan mengekstrak nilai DOI dari hasil yang paling cocok dengan artikel.",
      o="DOI yang berhasil diperoleh dari CrossRef.",
      j="Tujuan dari blok kode ini adalah mencari DOI cadangan ketika DOI tidak tersedia pada metadata agar identitas artikel tetap dapat dipastikan."),

 dict(t="validasi duplikasi data artikel",
      i="Nilai field artikel (DOI, URL artikel, URL PDF) dan isi database Supabase beserta final dataset sementara.",
      p="Fungsi _exists(field, value) memeriksa apakah suatu nilai sudah tersimpan pada tabel database. Bila ditemukan mengembalikan True, bila tidak ditemukan atau terjadi kesalahan mengembalikan False.",
      o="Status duplikat (True/False) yang menentukan apakah artikel disimpan.",
      j="Tujuan dari blok kode ini adalah mencegah artikel yang sama tersimpan lebih dari sekali pada database maupun dataset final."),

 dict(t="validasi PDF (1): penentuan URL PDF terbaik",
      i="DOI artikel serta URL PDF yang berasal dari Google Scholar.",
      p="Sistem mencoba memperoleh PDF resmi berdasarkan DOI melalui find_official_pdf_from_doi(); jika ditemukan URL tersebut dipilih sebagai sumber utama, bila tidak tersedia maka memakai URL PDF Google Scholar sebagai alternatif.",
      o="URL PDF terbaik yang akan digunakan untuk mengunduh artikel.",
      j="Tujuan dari blok kode ini adalah memilih sumber PDF yang paling resmi dan dapat diakses agar dokumen yang disimpan benar-benar berkualitas."),

 dict(t="validasi PDF (2): unduh dan validasi file",
      i="URL PDF yang telah dipilih pada tahap sebelumnya.",
      p="Sistem membuat nama file aman (safe_title) dan menentukan lokasi penyimpanan, lalu mengirim permintaan unduhan dengan HTTP headers yang menyerupai akses browser. Setelah file diterima, sistem memeriksa apakah file benar berformat PDF dan tidak corrupt.",
      o="File PDF yang valid beserta status scraping (pdf_downloaded atau metadata_only).",
      j="Tujuan dari blok kode ini adalah memastikan file PDF benar terunduh dan utuh sehingga artikel dapat disimpan dengan dokumen lengkap."),

 dict(t="kode cleansing text",
      i="Teks mentah judul dan abstrak artikel sebelum diproses.",
      p="Proses cleansing dilakukan berurutan: menghapus URL menggunakan ekspresi reguler (re.sub), mengonversi tanda hubung menjadi spasi agar kata majemuk seperti web-based terpisah, serta operasi pembersihan karakter lainnya.",
      o="Teks yang telah bersih dari elemen non-semantik dan siap ditokenisasi.",
      j="Tujuan dari blok kode ini adalah menghilangkan noise non-semantik (misal tautan web) sehingga teks hanya memuat informasi yang bermakna."),

 dict(t="tokenization text",
      i="Teks hasil cleansing sebagai masukan fungsi.",
      p="Sistem mengecek tipe masukan; bila bukan string mengembalikan list kosong. Bila valid, teks dipecah menjadi token berdasarkan spasi.",
      o="Daftar token (list of str) dari teks masukan.",
      j="Tujuan dari blok kode ini adalah memecah teks menjadi unit-unit token yang menjadi dasar proses pembobotan."),

 dict(t="stopword removal text",
      i="Daftar token hasil tokenization.",
      p="Sistem menghapus token yang termasuk daftar stopword (seperti yang, dan, di, the, of, an) serta token yang hanya memiliki satu karakter.",
      o="Token bermakna yang bebas dari kata umum yang tidak membedakan dokumen.",
      j="Tujuan dari blok kode ini adalah mengurangi dominasi kata umum non-diskriminatif agar pembobotan TF-IDF berfokus pada istilah penting."),

 dict(t="stemming text",
      i="Token hasil stopword removal, terutama token berbahasa Indonesia.",
      p="Sistem menerapkan stemming algoritma Nazief-Adriani (PySastrawi) untuk mengubah kata berimbuhan bahasa Indonesia menjadi bentuk dasar. Pada token bahasa Inggris proses ini tidak banyak mengubah bentuk kata.",
      o="Token dalam bentuk dasar (root word) sehingga variasi morfologis diseragamkan.",
      j="Tujuan dari blok kode ini adalah menyeragamkan bentuk kata agar term yang bermakna sama tidak dihitung terpisah pada matriks TF-IDF."),

 dict(t="kode pembentukan dokumen VSM",
      i="Kolom title dan abstract pada dataset artikel.",
      p="Fungsi preprocess_to_tokens() diterapkan terpisah pada kolom judul dan abstrak menghasilkan title_tokens dan abstract_tokens, kemudian keduanya digabung menjadi document_tokens dan dikonversi menjadi document_text.",
      o="Kolom document_text per artikel yang memuat teks gabungan judul dan abstrak.",
      j="Tujuan dari blok kode ini adalah membentuk representasi teks gabungan per dokumen sebagai masukan perhitungan TF-IDF dan konstruksi matriks VSM."),

 dict(t="kode perhitungan DF dan IDF",
      i="Kumpulan document_text dari seluruh 200 dokumen.",
      p="Sistem menghitung Document Frequency (DF) yaitu banyaknya dokumen yang memuat suatu term, lalu menghitung Inverse Document Frequency (IDF) untuk memberi bobot diskriminatif pada tiap term.",
      o="Nilai IDF untuk setiap term pada vocabulary (total 1.733 term unik).",
      j="Tujuan dari blok kode ini adalah menghitung IDF agar term yang muncul pada sedikit dokumen mendapatkan bobot lebih besar dan lebih diskriminatif."),

 dict(t="kode perhitungan TF-IDF",
      i="Frekuensi kata (term) pada setiap dokumen dan nilai IDF yang telah dihitung.",
      p="Sistem mengalikan term frequency (TF) dengan inverse document frequency (IDF) untuk memperoleh bobot gabungan setiap term pada setiap dokumen.",
      o="Matriks bobot TF-IDF untuk seluruh pasangan dokumen-term.",
      j="Tujuan dari blok kode ini adalah menghitung bobot term yang menggabungkan frekuensi lokal dan distribusi global agar term penting mendapat bobot optimal."),

 dict(t="kode pembentukan matriks VSM",
      i="Indeks dokumen (doc_to_index), indeks term (term_to_index), dan nilai bobot TF-IDF.",
      p="Setiap dokumen dipetakan ke indeks baris, setiap term ke indeks kolom, lalu nilai bobot TF-IDF disusun dalam format sparse csr_matrix agar efisien dalam penyimpanan.",
      o="Matriks Vector Space Model berukuran 200 x 1.733 dalam bentuk sparse.",
      j="Tujuan dari blok kode ini adalah membangun representasi dokumen-term yang efisien dan siap dipakai untuk perhitungan kemiripan query dengan dokumen."),

 dict(t="Implementasi preprocessing query dan pembentukan vektor query",
      i="Query yang diketik pengguna pada kolom pencarian.",
      p="Fungsi preprocess_query() memproses query dengan tahapan identik dokumen (cleansing, tokenisasi, hapus stopword, stemming), lalu build_query_vector() mengubah token hasil preprocessing menjadi vektor berbobot TF-IDF.",
      o="Vektor query berbobot TF-IDF yang berada dalam ruang vektor yang sama dengan dokumen.",
      j="Tujuan dari blok kode ini adalah menyiapkan query dalam representasi vektor yang sebanding dengan dokumen agar tingkat kesamaannya dapat dihitung."),

 dict(t="fungsi search Cosine Similarity",
      i="Vektor query yang telah dibentuk dan matriks VSM dokumen.",
      p="Fungsi search() menghitung skor cosine similarity antara vektor query dan seluruh baris matriks menggunakan cosine_similarity() dari scikit-learn, kemudian mengurutkan hasil dan mengembalikan Top-10 artikel dengan skor tertinggi.",
      o="Daftar artikel yang diurutkan berdasarkan nilai kemiripan (top-k).",
      j="Tujuan dari blok kode ini adalah memeringkat artikel berdasarkan kemiripan terhadap query sehingga dokumen paling relevan tampil paling atas hasil pencarian."),

 dict(t="implementasi fungsi manual Cosine",
      i="Dua vektor vector_a dan vector_b yang merepresentasikan dokumen.",
      p="Fungsi manual_cosine(vector_a, vector_b) menghitung dot product melalui np.dot, lalu hasilnya dibagi dengan produk norma kedua vektor menggunakan np.linalg.norm.",
      o="Nilai skor kemiripan kosinus antara 0 hingga 1 untuk kedua vektor.",
      j="Tujuan dari blok kode ini adalah memvalidasi hasil perhitungan cosine similarity manual dibandingkan keluaran scikit-learn agar akurasi metode terjamin."),

 dict(t="kode persiapan query uji",
      i="Daftar 12 query uji bilingual (Indonesia dan Inggris) beserta kategori masing-masing.",
      p="Variabel queries_eval mendefinisikan parameter evaluasi Information Retrieval. Setiap query dijalankan dengan top_k = 20, hasil dinilai oleh tiga evaluator independen, Label Final ditentukan melalui konsensus mayoritas, dan kesepakatan diukur dengan Fleiss Kappa.",
      o="Dataset evaluasi lengkap dengan ground truth (Label Final) beserta parameter K untuk menghitung Precision@K.",
      j="Tujuan dari blok kode ini adalah menyiapkan dataset evaluasi agar ground truth terdefinisi dengan jelas sehingga metrik Precision@K dapat dihitung secara valid."),

 dict(t="kode pencarian testing (Precision@K)",
      i="Seluruh query pada queries_eval, hasil pencarian sistem, dan data ground truth (evaluator_df).",
      p="Sistem mengiterasi seluruh query, mengambil pasangan query-artikel, mengonversi Label Final menjadi nilai biner, lalu menghitung Relevan@K dan P@K untuk setiap K (5, 10, 20). Hasil disimpan ke berkas CSV.",
      o="Tabel hasil berisi Relevan@K dan P@K untuk setiap query uji.",
      j="Tujuan dari blok kode ini adalah menghitung Precision@K untuk setiap query uji agar performa sistem Information Retrieval dapat dievaluasi terhadap ground truth."),

 dict(t="Halaman Eksplorasi pada sistem PaperCitation (UI Search)",
      i="Interaksi pengguna pada halaman awal: mengetikkan kata kunci dan memilih parameter pencarian.",
      p="Halaman menyediakan kolom input query, tombol pencarian, filter, rekomendasi kata kunci, riwayat pencarian, tips pencarian, serta fitur request artikel. Aksi pengguna diterjemahkan menjadi permintaan pencarian ke backend.",
      o="Request pencarian ke backend yang memuat query dan parameter yang dipilih.",
      j="Tujuan dari blok kode ini adalah menyediakan pintu masuk pengguna agar proses pencarian dapat dimulai secara intuitif tanpa perlu memahami teknis TF-IDF dan Cosine Similarity."),

 dict(t="kode program search (endpoint /search)",
      i="Parameter query beserta filter yang dikirim frontend melalui HTTP GET.",
      p="Endpoint /search memvalidasi bahwa query tidak kosong (mengembalikan HTTP 400 bila kosong), membaca seluruh parameter (top_k, tahun publikasi, jenis artikel, jenis analisis, jumlah kemunculan), lalu menjalankan pipeline pencarian.",
      o="Response JSON yang memuat daftar hasil pencarian yang deterministik sesuai query.",
      j="Tujuan dari blok kode ini adalah menghubungkan frontend dan backend agar query diproses dan hasil dikembalikan dalam format JSON."),

 dict(t="tampilan web proses scraping (progress dialog)",
      i="Status proses yang dikirim backend melalui Server-Sent Events (SSE).",
      p="Dialog menampilkan progress bar, daftar tahapan proses beserta statusnya, micro-phrases yang mendeskripsikan detail setiap tahap, serta animasi penulisan teks (typing effect).",
      o="Gambaran kemajuan proses scraping secara real-time kepada pengguna.",
      j="Tujuan dari blok kode ini adalah memberikan umpan balik real-time kepada pengguna agar status proses scraping yang sedang berjalan dapat dipantau."),

 dict(t="kode search-progress (endpoint /search-progress)",
      i="Keyword dan parameter pencarian dari frontend.",
      p="Endpoint memvalidasi query tidak kosong, membaca seluruh parameter, menjalankan pipeline pencarian/scraping dalam thread terpisah, lalu mengirimkan progress secara streaming melalui SSE.",
      o="Rangkaian event progress yang dikonsumsi frontend secara real-time.",
      j="Tujuan dari blok kode ini adalah menjalankan proses pencarian secara asinkron sekaligus memantau perkembangannya secara real-time kepada pengguna."),

 dict(t="tampilan web Filtering Search",
      i="Pilihan filter yang diambil pengguna pada halaman filter pencarian publikasi.",
      p="Pengguna menentukan parameter seperti jenis artikel, rentang tahun publikasi, kategori penelitian, jenis analisis, dan jumlah kemunculan kata kunci, kemudian menekan tombol Terapkan Filter.",
      o="Parameter filter yang siap dikirim ke backend untuk mempersempit hasil pencarian.",
      j="Tujuan dari blok kode ini adalah mempersempit hasil pencarian sesuai kebutuhan spesifik pengguna."),

 dict(t="kode Filtering Search",
      i="Filter yang telah dipilih (SearchFilters) beserta query.",
      p="Fungsi fetchWithFilters memvalidasi bahwa query tidak kosong, membaca filter (jenis artikel, jenis analisis, tahun, jumlah kemunculan), mengonversi nilai sitasi menjadi angka dengan validasi, lalu memanggil searchArticles untuk memperoleh hasil.",
      o="Daftar artikel hasil pencarian yang telah disaring sesuai filter.",
      j="Tujuan dari blok kode ini adalah menerapkan filter pada proses pencarian dan memperbarui hasil yang ditampilkan sesuai kriteria pengguna."),

 dict(t="tampilan Beranda Hasil Search",
      i="Data hasil pencarian yang dikembalikan backend berupa daftar artikel dan statistik.",
      p="Halaman menampilkan ringkasan jumlah artikel dan total kemunculan kata kunci, grafik batang distribusi publikasi berdasarkan tahun, serta tabel hasil yang diurutkan berdasarkan nilai kemiripan.",
      o="Tampilan hasil pencarian yang informatif (ringkasan, visualisasi, dan tabel).",
      j="Tujuan dari blok kode ini adalah menyajikan hasil pencarian agar pengguna dapat membaca distribusi dan himpunan artikel relevan dengan mudah."),

 dict(t="kode Beranda Hasil Search",
      i="Daftar artikel hasil pencarian beserta statistik (paperCount, totalMatched, trendData).",
      p="Komponen SectionCards menampilkan informasi inti hasil pencarian, ChartBarLabel merender grafik distribusi tahun, dan DataTable menampilkan tabel artikel. Seluruh komponen disusun dalam satu halaman.",
      o="Halaman hasil lengkap berupa ringkasan, visualisasi, dan tabel artikel.",
      j="Tujuan dari blok kode ini adalah merender hasil pencarian dalam satu kesatuan ringkasan, grafik, dan tabel agar informasi mudah dicerna pengguna."),

 dict(t="tampilan daftar artikel",
      i="Seluruh artikel yang tersimpan pada basis data.",
      p="Halaman menampilkan tabel yang memuat judul, penulis, tahun, jenis artikel, kategori, dan aksi pengelolaan; pengguna dapat melakukan sorting berdasarkan relevansi atau tahun.",
      o="Daftar artikel yang ditampilkan dalam bentuk tabel.",
      j="Tujuan dari blok kode ini adalah menyediakan media pengelolaan data artikel sekaligus menjadi sumber data utama proses pencarian publikasi."),

 dict(t="kode daftar artikel",
      i="Query pencarian dan filter aktif yang tersimpan pengguna.",
      p="Fungsi fetchArticles memvalidasi query (mengalihkan ke halaman /search bila kosong), mengaktifkan status loading, membaca filter tersimpan, memanggil API untuk mengambil data, lalu mengurutkan hasil sesuai opsi pengurutan.",
      o="Tabel berisi daftar artikel yang diambil dan diurutkan sesuai kriteria.",
      j="Tujuan dari blok kode ini adalah mengambil, mengolah, dan menampilkan data artikel hasil pencarian sesuai urutan yang diinginkan pengguna."),
]

# ===========================================================================
# FUNGSI BANTU
# ===========================================================================
def add_h(doc, text, level):
    doc.add_heading(text, level=level)


def add_p(doc, text, bold=False, italic=False, color=None, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = RGBColor(*color)
    if size is not None:
        r.font.size = Pt(size)
    return p


def add_label(doc, label, text):
    p = doc.add_paragraph()
    r = p.add_run(label + ": ")
    r.bold = True
    p.add_run(text)
    return p


def add_catatan(doc, text):
    p = doc.add_paragraph()
    r = p.add_run("[CATATAN] ")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p.add_run(text)
    return p


# ===========================================================================
# DOKUMEN 1: Revisi_Narasi_5Poin.docx
# ===========================================================================
def build_doc1():
    doc = Document()

    add_p(doc, "REVISI NARASI LAPORAN AKHIR - 5 POIN REVISI DOSEN PEMBIMBING",
          bold=True, size=16)
    add_p(doc, "Panduan pemakaian: teks berpenanda [SISIPKAN] adalah kalimat baru untuk disisipkan; "
               "[GANTI] adalah pengganti kalimat lama; nomor tabel/gambar memakai penomoran otomatis "
               "SEQ di Word.", italic=True)

    # ---- BAGIAN 1 ----
    add_h(doc, "BAGIAN 1 - Penjelasan Kode dengan Pendekatan Sandwich (Input-Proses-Output)", 1)
    add_p(doc, "Gunakan pendekatan Sandwich Explanatory untuk setiap gambar kode yang krusial "
               "(Gambar 4.1 s.d. 4.33). Setiap blok dijelaskan dengan tiga komponen: INPUT (data yang "
               "masuk ke fungsi), PROSES (logika inti tanpa menulis ulang sintaks kode), dan OUTPUT "
               "(data yang dihasilkan), lalu ditutup dengan kalimat tujuan.")
    for i, f in enumerate(F, start=1):
        add_h(doc, "Gambar 4.%d - %s" % (i, f["t"]), 2)
        add_label(doc, "INPUT", f["i"])
        add_label(doc, "PROSES", f["p"])
        add_label(doc, "OUTPUT", f["o"])
        add_p(doc, f["j"], italic=True)
    add_h(doc, "Contoh kalimat penutup seragam (sesuai permintaan dosen)", 2)
    add_p(doc, "Tujuan dari blok kode ini adalah untuk [tujuan] agar menghasilkan [output], yang "
               "kemudian akan diproses ke tahap [tahap selanjutnya].")

    # ---- BAGIAN 2 ----
    add_h(doc, "BAGIAN 2 - Justifikasi Kriteria Interpretasi (Nilai Threshold)", 1)
    add_catatan(doc, "Sisipkan paragraf di bawah ini pada Bab IV di sekitar Tabel 4.33 (Hasil "
                     "Eksperimen Penentuan Threshold), sebelum atau sesudah tabel, agar pemilihan "
                     "rentang threshold dikaitkan dengan urgensi masalah.")
    add_h(doc, "[SISIPKAN] Argumen Justifikasi Threshold", 2)
    add_p(doc, "Penentuan nilai threshold pada rentang 0,7-0,95 sebagai kriteria interpretasi "
               "didasarkan pada urgensi kebutuhan presisi tinggi dalam pencarian literatur akademik, "
               "di mana tingkat toleransi kesalahan (error rate) harus diminimalisir agar peneliti "
               "tidak terdistraksi oleh dokumen yang tidak relevan. Pada penelitian ini, rentang "
               "tersebut digunakan untuk menginterpretasikan kategori kemiripan hasil cosine "
               "similarity, yaitu Relevan Tinggi (>= 0,70), Relevan Sedang (0,40-0,70), dan Rendah "
               "(< 0,40), mengikuti kriteria interpretasi Sudiatmika et al. (2026).")
    add_p(doc, "Adapun threshold 0,30 yang digunakan sebagai ambang operasional filter sistem "
               "diperoleh dari hasil eksperimen terhadap data ground truth (Tabel 4.33), sehingga "
               "kedua peran threshold dapat dibedakan dengan jelas: rentang 0,7-0,95 berfungsi "
               "sebagai kriteria interpretasi kategori, sedangkan 0,30 berfungsi sebagai ambang "
               "operasional yang menyeimbangkan precision dan recall pada proses retrieval.")

    # ---- BAGIAN 3 ----
    add_h(doc, "BAGIAN 3 - Bukti Efektivitas Metode (Analisis Komparatif Eksperimen)", 1)
    add_h(doc, "[SISIPKAN] di Bab 4.2 Pembahasan - Analisis Komparatif Tabel 4.33", 2)
    add_p(doc, "Berdasarkan Tabel 4.33, terlihat bahwa precision tetap konstan sebesar 1,00 pada "
               "seluruh nilai threshold (0,30 sampai 0,70), artinya sistem tidak pernah menghasilkan "
               "false positive pada semua ambang. Namun, seiring kenaikan threshold, recall menurun "
               "drastis: 0,4490 (threshold 0,30), 0,2041 (0,40), 0,0969 (0,50), 0,0204 (0,60), dan "
               "0,0051 (0,70). Hal ini berarti semakin tinggi ambang, semakin banyak artikel yang "
               "sebenarnya relevan berdasarkan human judgment terlewat oleh sistem (false negative "
               "naik dari 108 menjadi 195).")
    add_p(doc, "Fenomena tersebut membuktikan bahwa metode VSM dengan cosine similarity pada sistem "
               "ini bekerja optimal pada ambang threshold 0,30, yaitu nilai yang menyeimbangkan "
               "precision (1,00) dan recall (0,4490) sehingga sistem tetap mampu mengambil artikel "
               "relevan secara maksimal tanpa menghasilkan noise berupa artikel tidak relevan. Pada "
               "threshold tinggi, precision tinggi menjadi tidak berguna karena recall praktis mendekati "
               "nol, sehingga dokumen yang relevan hampir tidak terambil.")
    add_p(doc, "Oleh karena itu, efektivitas metode tidak hanya diukur dari tingginya precision, "
               "melainkan dari kesetimbangan precision dan recall. Hasil eksperimen ini memperkuat "
               "argumentasi bahwa tahapan preprocessing pada metadata artikel ilmiah berkontribusi "
               "meningkatkan kualitas kemiripan, sehingga ambang yang rendah sudah mampu menghasilkan "
               "precision penuh tanpa false positive.")

    # ---- BAGIAN 4 ----
    add_h(doc, "BAGIAN 4 - Validitas Data Ground Truth (Bab 3.3.6 Metode Pengujian)", 1)
    add_catatan(doc, "Tambahkan poin berikut pada sub-bab 3.3.6 Metode Pengujian untuk menegaskan "
                     "bagaimana ground truth ditentukan dan divalidasi.")
    add_label(doc, "Siapa yang memvalidasi",
              "Data dinilai oleh tiga evaluator independen yang merupakan mahasiswa Politeknik "
              "Negeri Banyuwangi melalui blind evaluation, yaitu menilai tanpa mengetahui skor "
              "similarity yang dihasilkan sistem. Evaluator dibekali panduan labeling (Lampiran 8) "
              "untuk menjaga konsistensi penilaian.")
    add_label(doc, "Kriteria relevansi",
              "Setiap pasangan query-artikel diberi label biner (1 = relevan, 0 = tidak relevan) "
              "berdasarkan kesesuaian judul dan abstrak artikel terhadap query. Label Final "
              "ditentukan melalui konsensus mayoritas ketiga evaluator dan dijadikan ground truth "
              "pada perhitungan Precision@K.")
    add_label(doc, "Pengendalian kualitas penilaian",
              "Kesepakatan antar evaluator diukur menggunakan Fleiss Kappa pada 224 pasangan "
              "query-artikel dengan nilai kappa = 0,176 (kategori slight agreement). Meskipun "
              "tergolong rendah karena subjektivitas dalam menilai relevansi, penggunaan konsensus "
              "mayoritas memastikan satu Label Final yang konsisten sebagai pembanding evaluasi.")
    add_label(doc, "Proses data",
              "Hasil pencarian Top-K dari sistem dicocokkan terhadap ground truth untuk menghitung "
              "Relevan@K dan P@K dengan K = 5, 10, dan 20 pada 12 query bilingual yang mewakili "
              "empat kategori.")

    # ---- BAGIAN 5 ----
    add_h(doc, "BAGIAN 5 - Perbandingan dengan Penelitian Terdahulu (Bab 4.2 Pembahasan)", 1)
    add_catatan(doc, "Sisipkan paragraf berikut di akhir Bab 4.2 Pembahasan, setelah paragraf yang "
                     "menyebut Heryawan et al. (2024) dan Nurkholis et al. (2023).")
    add_h(doc, "[SISIPKAN] Analisis Kontribusi terhadap Penelitian Terdahulu", 2)
    add_p(doc, "Dibandingkan dengan hasil penelitian terdahulu, sistem yang dikembangkan dalam "
               "penelitian ini menunjukkan peningkatan kinerja yang signifikan. Nurkholis et al. "
               "(2023) melaporkan nilai precision sebesar 56% dan recall 98% saat menerapkan VSM "
               "untuk web scraping pada website freelance, sedangkan Heryawan et al. (2024) "
               "memperoleh precision 54,8% dan recall 79,6% pada pencarian dokumen rekam medis. "
               "Sebagai perbandingan, sistem ini menghasilkan rata-rata Precision@5 sebesar 90,00%, "
               "Precision@10 sebesar 86,67%, dan Precision@20 sebesar 81,67%.")
    add_p(doc, "Peningkatan ini diindikasikan karena adanya optimasi pada tahap preprocessing yang "
               "lebih spesifik untuk metadata artikel ilmiah, mulai dari similarity checking, "
               "pemilihan sumber resmi, validasi DOI dan PDF, hingga preprocessing teks (cleansing, "
               "tokenisasi, stopword removal, stemming) yang disesuaikan dengan karakteristik bahasa "
               "Indonesia. Dengan demikian, pendekatan yang diajukan terbukti memberikan kontribusi "
               "yang lebih kompetitif dibandingkan penelitian sebelumnya pada domain pencarian "
               "artikel ilmiah bilingual.")

    out = os.path.join(OUT_DIR, "Revisi_Narasi_5Poin.docx")
    doc.save(out)
    return out


# ===========================================================================
# DOKUMEN 2: BabIV_Gabungan_Revisi.docx
# ===========================================================================
GROUPS = [
    ("4.1.1", "Implementasi Pipeline Pengambilan dan Validasi Data Artikel", 0, 10,
     "Bagian ini menjelaskan implementasi pipeline pengambilan data artikel dari Google Scholar "
     "beserta seluruh tahap validasi agar kualitas data terjamin, mulai dari snapshot metadata, "
     "pre-filtering, similarity checking, pemilihan sumber resmi, validasi DOI dan duplikasi, "
     "hingga validasi PDF."),
    ("4.1.2", "Preprocessing Teks", 10, 4,
     "Agar teks dapat dihitung kemiripannya, seluruh judul dan abstrak artikel melewati empat "
     "tahap: cleansing, tokenisasi, penghapusan stopword, dan stemming."),
    ("4.1.3", "Pembentukan Vektor dengan Pembobotan TF-IDF", 14, 4,
     "Setelah pra-pemrosesan, sistem membangun representasi dokumen berupa vektor menggunakan "
     "metode TF-IDF dan menyusunnya ke dalam matriks Vector Space Model (VSM)."),
    ("4.1.4", "Proses Pencarian dengan Cosine Similarity", 18, 3,
     "Bagian ini menjelaskan bagaimana query pengguna diubah menjadi vektor dan dibandingkan "
     "dengan seluruh dokumen untuk menghasilkan peringkat artikel berdasarkan tingkat kemiripan."),
    ("4.1.5", "Pengujian Precision at K", 21, 2,
     "Evaluasi sistem dilakukan menggunakan metrik Precision@K terhadap data ground truth yang "
     "diperoleh dari penilaian tiga evaluator independen."),
    ("4.1.6", "Implementasi Antarmuka Sistem", 23, 10,
     "Antarmuka sistem berfungsi sebagai penghubung antara pengguna dan proses informatif "
     "retrieval di sisi backend."),
]


def build_doc2():
    doc = Document()
    add_h(doc, "BAB IV - HASIL DAN PEMBAHASAN", 0)
    add_p(doc, "Bab ini menyajikan hasil implementasi dan pengujian sistem pencarian artikel "
               "ilmiah berbasis web scraping dengan Vector Space Model, pembobotan TF-IDF, dan "
               "Cosine Similarity, dilanjutkan dengan pembahasan untuk menjawab rumusan masalah.")

    add_h(doc, "4.1 Hasil Penelitian", 1)
    for no, judul, start, n, intro in GROUPS:
        add_h(doc, "%s %s" % (no, judul), 2)
        add_p(doc, intro)
        for k in range(start, start + n):
            f = F[k]
            add_h(doc, "Gambar 4.%d - %s" % (k + 1, f["t"]), 3)
            add_label(doc, "INPUT", f["i"])
            add_label(doc, "PROSES", f["p"])
            add_label(doc, "OUTPUT", f["o"])
            add_p(doc, f["j"], italic=True)

    # ---- 4.1.7 Eksperimen Threshold ----
    add_h(doc, "4.1.7 Eksperimen Penentuan Threshold Cosine Similarity", 2)
    add_p(doc, "Untuk memvalidasi nilai threshold yang digunakan pada sistem, dilakukan eksperimen "
               "penentuan threshold dengan membandingkan prediksi kelas sistem terhadap label ground "
               "truth. Hasil eksperimen ditunjukkan pada Tabel 4.33.")
    add_p(doc, "Tabel 4.33. Hasil Eksperimen Penentuan Threshold Cosine Similarity", bold=True)
    add_p(doc, "Threshold: 0,30 | 0,40 | 0,50 | 0,60 | 0,70")
    add_p(doc, "Accuracy: 0,5200 | 0,3067 | 0,2133 | 0,1467 | 0,1333")
    add_p(doc, "Precision: 1,0000 | 1,0000 | 1,0000 | 1,0000 | 1,0000")
    add_p(doc, "Recall: 0,4490 | 0,2041 | 0,0969 | 0,0204 | 0,0051")
    add_p(doc, "F1-Score: 0,6197 | 0,3390 | 0,1767 | 0,0400 | 0,0102")
    add_p(doc, "Penentuan nilai threshold pada rentang 0,7-0,95 sebagai kriteria interpretasi "
               "didasarkan pada urgensi kebutuhan presisi tinggi dalam pencarian literatur "
               "akademik, di mana tingkat toleransi kesalahan (error rate) harus diminimalisir agar "
               "peneliti tidak terdistraksi oleh dokumen yang tidak relevan. Rentang tersebut "
               "digunakan untuk menginterpretasikan kategori kemiripan hasil cosine similarity, "
               "yaitu Relevan Tinggi (>= 0,70), Relevan Sedang (0,40-0,70), dan Rendah (< 0,40), "
               "mengikuti Sudiatmika et al. (2026).")
    add_p(doc, "Adapun threshold 0,30 sebagai ambang operasional filter diperoleh dari hasil "
               "eksperimen pada Tabel 4.33, sehingga keduanya dapat dibedakan: rentang 0,7-0,95 "
               "adalah kriteria interpretasi kategori, sedangkan 0,30 adalah ambang operasional "
               "yang menyeimbangkan precision dan recall pada proses retrieval.")

    # ---- 4.1.8 Hasil Pengujian ----
    add_h(doc, "4.1.8 Hasil Pengujian Precision at K", 2)
    add_p(doc, "Hasil pengujian diperoleh dari proses pencarian terhadap 12 query uji yang mewakili "
               "empat kategori. Setiap query dicocokkan dengan ground truth untuk menghitung nilai "
               "precision pada K=5, 10, dan 20. Sistem menghasilkan Mean P@5 sebesar 90,00%, Mean "
               "P@10 sebesar 86,67%, dan Mean P@20 sebesar 81,67%. Penurunan nilai precision "
               "seiring bertambahnya nilai K merupakan pola umum pada sistem Information Retrieval "
               "karena semakin banyak dokumen yang ditampilkan, peluang munculnya dokumen tidak "
               "relevan semakin meningkat.")

    # ---- 4.2 Pembahasan ----
    add_h(doc, "4.2 Pembahasan", 1)
    add_p(doc, "Berdasarkan hasil penelitian, proses web scraping berhasil mengumpulkan 200 artikel "
               "ilmiah dari Google Scholar yang terbagi merata ke dalam empat kategori. Seluruh "
               "data berhasil diproses melalui tahapan preprocessing sehingga menghasilkan dokumen "
               "yang siap digunakan pada proses pembobotan TF-IDF.")
    add_p(doc, "Hasil pembobotan TF-IDF membentuk vocabulary sebanyak 1.733 term dengan matriks "
               "Vector Space Model berukuran 200 x 1.733 dan tingkat sparsity sebesar 98,88%. "
               "Kondisi ini merupakan karakteristik umum representasi teks berbasis TF-IDF yang "
               "menghasilkan vektor bersifat sparse, sesuai teori Information Retrieval (Azizah & "
               "Handayani, 2022).")

    add_h(doc, "Analisis Efektivitas Metode berdasarkan Eksperimen Threshold", 2)
    add_p(doc, "Berdasarkan Tabel 4.33, precision tetap konstan sebesar 1,00 pada seluruh nilai "
               "threshold (0,30 sampai 0,70), artinya sistem tidak pernah menghasilkan false "
               "positive pada semua ambang. Namun, seiring kenaikan threshold, recall menurun "
               "drastis dari 0,4490 menjadi 0,0051, sehingga semakin banyak artikel yang relevan "
               "berdasarkan human judgment terlewat oleh sistem (false negative naik dari 108 "
               "menjadi 195). Hal ini membuktikan bahwa metode VSM dengan cosine similarity bekerja "
               "optimal pada ambang threshold 0,30, yaitu nilai yang menyeimbangkan precision "
               "(1,00) dan recall (0,4490) tanpa menghasilkan noise.")
    add_p(doc, "Efektivitas metode tidak hanya diukur dari tingginya nilai precision, melainkan "
               "dari kesetimbangan antara precision dan recall. Hasil eksperimen ini memperkuat "
               "argumentasi bahwa tahapan preprocessing pada metadata artikel ilmiah berkontribusi "
               "meningkatkan kualitas kemiripan sehingga ambang rendah sudah mampu menghasilkan "
               "precision penuh tanpa false positive.")

    add_h(doc, "Analisis Performa per Kategori", 2)
    add_p(doc, "Kategori Cyber Security memiliki performa paling stabil karena penggunaan istilah "
               "yang relatif homogen, sedangkan kategori Machine Learning mengalami penurunan "
               "precision pada query berbahasa Indonesia. Hal ini menunjukkan adanya tantangan pada "
               "pencarian bilingual, yaitu perbedaan frekuensi penggunaan istilah Indonesia dan "
               "Inggris di dalam korpus yang memengaruhi tingkat kemiripan dokumen terhadap query.")

    add_h(doc, "Perbandingan dengan Penelitian Terdahulu", 2)
    add_p(doc, "Hasil penelitian ini mendukung Heryawan et al. (2024) dan Nurkholis et al. (2023) "
               "yang menyatakan bahwa kombinasi TF-IDF dan Cosine Similarity efektif meningkatkan "
               "relevansi hasil pencarian. Dibandingkan Nurkholis et al. (2023) dengan precision "
               "56% serta Heryawan et al. (2024) dengan precision 54,8%, sistem ini menghasilkan "
               "rata-rata Precision@5 sebesar 90,00% dan Precision@10 sebesar 86,67%. Peningkatan "
               "ini diindikasikan karena adanya optimasi pada tahap preprocessing yang lebih "
               "spesifik untuk metadata artikel ilmiah, sehingga pendekatan yang diajukan "
               "memberikan kontribusi yang lebih kompetitif pada pencarian artikel ilmiah evaluasi "
               "precision and recall.")

    add_p(doc, "Secara keseluruhan, hasil penelitian menunjukkan bahwa sistem berhasil "
               "mengimplementasikan web search scraping, Vector Space Model dengan TF-IDF dan "
               "Cosine Similarity, serta menghasilkan evaluasi yang baik melalui Precision@K. "
               "Keterbatasan penelitian mencakup jumlah dataset yang terbatas, penggunaan "
               "PySastrawi yang optimal hanya untuk bahasa Indonesia, dan keterbatasan TF-IDF "
               "dalam menangkap hubungan semantik antar kata.")

    out = os.path.join(OUT_DIR, "BabIV_Gabungan_Revisi.docx")
    doc.save(out)
    return out


# ===========================================================================
# MAIN
# ===========================================================================
if __name__ == "__main__":
    p1 = build_doc1()
    p2 = build_doc2()
    print("OK:", p1)
    print("OK:", p2)

